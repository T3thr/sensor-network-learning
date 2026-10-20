from pathlib import Path
import shutil, json, hashlib, re
from zipfile import ZipFile
from lxml import etree as ET
from pypdf import PdfReader
from docx import Document

ROOT = Path(__file__).resolve().parent.parent
BUILD = ROOT / '.build_v4'
report = 'Assignment_Paper_4_PublishTimeOpt_Short_Report_TH'
slides = 'Assignment_Paper_4_PublishTimeOpt_7_Slides'
shutil.copyfile(BUILD / 'report-render-research' / (report + '.pdf'), ROOT / '02_report' / (report + '.pdf'))
shutil.copyfile(BUILD / 'slides-pdf-research' / (slides + '.pdf'), ROOT / '03_slides' / (slides + '.pdf'))
fonts = ROOT / '03_slides/fonts'
fonts.mkdir(exist_ok=True)
shutil.copyfile(BUILD / 'IBMPlexThai-Regular.ttf', fonts / 'IBMPlexSansThai-Regular.ttf')
shutil.copyfile(BUILD / 'IBMPlexThai-Bold.ttf', fonts / 'IBMPlexSansThai-Bold.ttf')
shutil.copyfile(BUILD / 'IBMPlexSansThai-OFL.txt', fonts / 'OFL.txt')

metrics = json.loads((ROOT / '05_study/paper-metrics.json').read_text())
doc = Document(ROOT / '02_report' / (report + '.docx'))
headings = [p for p in doc.paragraphs if p.style is not None and p.style.style_id == 'Heading1']
assert len(headings) == 8
assert all(not re.match(r'^\s*\d', p.text) for p in headings)
assert all(p._p.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pPr/{http://schemas.openxmlformats.org/wordprocessingml/2006/main}numPr') is None for p in headings)
assert all(p.runs and all(r.bold for r in p.runs if r.text.strip()) for p in headings)
table = doc.tables[0]
def pair(text):
    cleaned = text.replace('−', '-').replace('–', '-')
    values = re.findall(r'[+-]?\d+\.\d+', cleaned)
    return tuple(float(v) for v in values)
expected = [[(r[m]['mean'], r[m]['sd']) for m in ['sc', 'ttr', 'tta']] for r in metrics['table_1']['rows']]
observed = [[pair(c.text) for c in row.cells[1:]] for row in table.rows[1:]]
assert observed == expected, (observed, expected)

ns = {'p': 'http://schemas.openxmlformats.org/presentationml/2006/main', 'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}
personal_voice = re.compile(r'ผม|ดิฉัน|ฉัน|พวกเรา|พวกคุณ|ของเรา|ของคุณ|ที่เรา|ที่คุณ|(?:^|\s)(?:เรา|คุณ)(?:\s|$)')
assert all(not personal_voice.search(p.text) for p in doc.paragraphs)
with ZipFile(ROOT / '03_slides' / (slides + '.pptx')) as z:
    for i in range(1, 8):
        xml = ET.fromstring(z.read(f'ppt/slides/slide{i}.xml'))
        slide_text = '\n'.join(xml.xpath('//a:t/text()', namespaces=ns))
        assert not personal_voice.search(slide_text)
    s6 = ET.fromstring(z.read('ppt/slides/slide6.xml'))
    s6_text = '\n'.join(s6.xpath('//a:t/text()', namespaces=ns))
    assert 'ผลงานวิจัยและข้อจำกัดของการทดลอง' in s6_text
    assert 'ผลงานสำคัญของบทความ' in s6_text and 'ข้อจำกัดของการทดลอง' in s6_text
    last = ET.fromstring(z.read('ppt/slides/slide7.xml'))
    last_text = '\n'.join(last.xpath('//a:t/text()', namespaces=ns))
    assert 'สรุปผลและแนวทางวิจัยในอนาคต' in last_text
    assert 'ผลที่ได้จากการทดลอง' in last_text and 'แนวทางที่ผู้วิจัยเสนอ' in last_text
    assert 'การอ่าน' not in s6_text + last_text
    for i in range(2, 8):
        page_xml = ET.fromstring(z.read(f'ppt/slides/slide{i}.xml'))
        footers = [shape for shape in page_xml.findall('.//p:sp', ns)
                   if shape.find('.//a:off', ns) is not None
                   and shape.find('.//a:off', ns).get('x') == '12954000'
                   and shape.find('.//a:off', ns).get('y') == '7953375']
        assert len(footers) == 1, (i, len(footers))
        assert ''.join(footers[0].xpath('.//a:t/text()', namespaces=ns)) == str(i-1)
    assert all(term in last_text for term in ['Hydration', 'Incremental Regeneration', 'Hybrid'])
    assert 'divwy' not in last_text.lower() and 'โครงงาน' not in last_text
    sx = ET.fromstring(z.read('ppt/slides/slide5.xml'))
    native = sx.find('.//a:tbl', ns)
    assert native is not None
    data = [[pair(''.join(c.xpath('.//a:t/text()', namespaces=ns))) for c in row.findall('a:tc', ns)[1:]] for row in native.findall('a:tr', ns)[1:]]
    assert data == expected, (data, expected)
    notes = []
    for i in range(1, 8):
        nx = ET.fromstring(z.read(f'ppt/notesSlides/notesSlide{i}.xml'))
        text = '\n'.join(nx.xpath('//a:t/text()', namespaces=ns))
        assert 'Source:' in text and metrics['doi'] in text
        assert not personal_voice.search(text)
        notes.append({'slide': i, 'source_and_doi_present': True, 'character_count': len(text)})
    size = ET.fromstring(z.read('ppt/presentation.xml')).find('p:sldSz', ns)
    assert (size.get('cx'), size.get('cy')) == ('14630400', '8229600')
    assert len([n for n in z.namelist() if re.fullmatch(r'ppt/slides/slide\d+\.xml', n)]) == 7

report_pdf = PdfReader(ROOT / '02_report' / (report + '.pdf'))
slides_pdf = PdfReader(ROOT / '03_slides' / (slides + '.pdf'))
assert len(report_pdf.pages) == 2
assert len(slides_pdf.pages) == 7
for p in report_pdf.pages:
    assert abs(float(p.mediabox.width) - 595.28) < 1
    assert abs(float(p.mediabox.height) - 841.89) < 1
for p in slides_pdf.pages:
    assert abs(float(p.mediabox.width) / float(p.mediabox.height) - 16/9) < .001
source = ROOT / '01_original/Paper_IV_1_PRIMARY_PublishTimeOpt_11298742.pdf'
assert hashlib.sha256(source.read_bytes()).hexdigest() == metrics['source_sha256']
reference = ROOT.parent / 'work3/02_report_v3/Assignment_Paper_3_v3_Short_Report_TH.docx'
assert hashlib.sha256(reference.read_bytes()).hexdigest() == 'b635f2a4ec31652bba572f1d1ca71e7c85d52790c8392d04eda2a1266182714c'
outputs = {}
for folder, stem, ext in [('02_report', report, 'docx'), ('02_report', report, 'pdf'), ('03_slides', slides, 'pptx'), ('03_slides', slides, 'pdf')]:
    f = ROOT / folder / (stem + '.' + ext)
    outputs[str(f.relative_to(ROOT))] = {'bytes': f.stat().st_size, 'sha256': hashlib.sha256(f.read_bytes()).hexdigest()}
result = {
    'status': 'pass', 'report_pdf_pages': 2, 'report_size': 'A4',
    'slide_count': 7, 'slide_ratio': '16:9',
    'report_headings_bold_without_numbering': True,
    'personal_pronouns_absent_from_report_slides_and_notes': True,
    'closing_slide': {'experimental_findings': 'ผลที่ได้จากการทดลอง', 'authors_future_work': 'แนวทางที่ผู้วิจัยเสนอ', 'paper_fidelity': True},
    'content_page_numbers': [1, 2, 3, 4, 5, 6],
    'study_scope': 'Paper architecture, experiments, limitations, and reader synthesis.',
    'report_table_pairs_verified_against_source_metrics': 9,
    'native_slide_table_pairs_verified_against_source_metrics': 9,
    'speaker_notes': notes, 'source_pdf_unchanged': True, 'reference_docx_unchanged': True,
    'visual_review': {'docx_render_pages': [1, 2], 'artifact_slides': list(range(1,8)), 'native_pdf_slides': list(range(1,8)), 'findings_remaining': []},
    'validation_limits': ['PDF rendering verified with bundled LibreOffice; Microsoft PowerPoint execution not tested.', '4:55 timing is an allocation, not a measured rehearsal.', 'TH SarabunPSK PDF glyph mapping limits Thai text extraction; visible report text and DOCX were checked.'],
    'outputs': outputs
}
(ROOT / '06_evidence').mkdir(exist_ok=True)
(ROOT / '06_evidence/report-slides-verification.json').write_text(json.dumps(result, ensure_ascii=False, indent=2))
print(json.dumps({k: v for k, v in result.items() if k not in ['speaker_notes', 'outputs']}, ensure_ascii=False, indent=2))
