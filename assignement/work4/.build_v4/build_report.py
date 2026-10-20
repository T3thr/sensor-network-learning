from pathlib import Path
from copy import deepcopy
import re, hashlib, json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
REFERENCE = ROOT.parent / 'work3/02_report_v3/Assignment_Paper_3_v3_Short_Report_TH.docx'
SOURCE = ROOT / '02_report/Assignment_Paper_4_PublishTimeOpt_Short_Report_TH.md'
OUT = SOURCE.with_suffix('.docx')
FONT = 'TH SarabunPSK'
NAVY = '0F2D4A'
INK = '0F172A'

doc = Document(REFERENCE)
body = doc._element.body
for element in list(body):
    if element.tag != qn('w:sectPr'):
        body.remove(element)

# Reuse the reference package and page furniture, authoring only the new content.
if 'Normal' not in doc.styles:
    doc.styles.add_style('Normal', WD_STYLE_TYPE.PARAGRAPH)
roles = {name: next(st for st in doc.styles if st.style_id == ident) for name, ident in [('Normal','Normal'),('Title','Title'),('Heading 1','Heading1')]}
for style_name, size in [('Normal', 16), ('Title', 22), ('Heading 1', 17)]:
    st = roles[style_name]
    st.font.name = FONT
    st.font.size = Pt(size)
    st.font.color.rgb = RGBColor.from_string(NAVY if style_name != 'Normal' else INK)
    st.paragraph_format.line_spacing = Pt(size * 1.15)
    st.paragraph_format.space_after = Pt(3)
    st.paragraph_format.space_before = Pt(0 if style_name != 'Heading 1' else 7)
    st.paragraph_format.keep_with_next = style_name != 'Normal'
    st.paragraph_format.widow_control = True
    if style_name != 'Normal': st.font.bold = True

def font_run(run, size, bold=False, color=INK):
    run.font.name = FONT; run.font.size = Pt(size); run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    rp = run._element.get_or_add_rPr()
    fonts = rp.find(qn('w:rFonts'))
    for key in ['ascii', 'hAnsi', 'cs', 'eastAsia']: fonts.set(qn('w:'+key),FONT)
    cs = OxmlElement('w:szCs'); cs.set(qn('w:val'),str(int(size*2))); rp.append(cs)
    if bold: rp.append(OxmlElement('w:bCs'))
    lang = OxmlElement('w:lang'); lang.set(qn('w:val'),'th-TH'); rp.append(lang)

def inline(p, text, size=16, color=INK):
    for token in re.split(r'(\*\*.*?\*\*)',text):
        bold = token.startswith('**') and token.endswith('**')
        text = token[2:-2] if bold else token
        text = re.sub(r'`([^`]+)`',r'\1',text)
        text = re.sub(r'(?<!\*)\*([^*]+)\*(?!\*)',r'\1',text)
        if text: font_run(p.add_run(text),size,bold,color)

def paragraph(text, size=16, meta=False):
    p=doc.add_paragraph()
    p.paragraph_format.space_after=Pt(3)
    p.paragraph_format.line_spacing=Pt(size * 1.15)
    p.paragraph_format.widow_control=True
    p.paragraph_format.keep_with_next=False
    p.alignment=WD_ALIGN_PARAGRAPH.LEFT if meta else WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent=Inches(0 if meta or re.match(r'^\d\.',text) else .3)
    inline(p,text,size)
    return p

def table(lines):
    vals=[[x.strip() for x in l.strip().strip('|').split('|')] for l in lines]
    vals=[v for v in vals if not all(re.fullmatch(r'[:\- ]+',x) for x in v)]
    t=doc.add_table(rows=len(vals), cols=4)
    t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    widths=[1.9,1.46,1.46,1.45]
    for i, row in enumerate(t.rows):
        props=row._tr.get_or_add_trPr()
        props.append(OxmlElement('w:cantSplit'))
        if i==0: props.append(OxmlElement('w:tblHeader'))
        for j, c in enumerate(row.cells):
            c.width=Inches(widths[j]); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=c._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'CBD5E1'); borders.append(e)
            pr.append(borders)
            sh=OxmlElement('w:shd'); sh.set(qn('w:fill'), NAVY if i==0 else ('F1F5F9' if i%2==0 else 'FFFFFF')); pr.append(sh)
            margins=OxmlElement('w:tcMar')
            for edge,value in [('top',65),('bottom',65),('left',100),('right',100)]:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),str(value));e.set(qn('w:type'),'dxa');margins.append(e)
            pr.append(margins)
            p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=Pt(16.5)
            p.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.CENTER
            font_run(p.add_run(vals[i][j]),14.5,i==0,'FFFFFF' if i==0 else INK)
    return t

lines=SOURCE.read_text().splitlines(); i=0
while i<len(lines):
    l=lines[i].strip()
    if not l: i+=1; continue
    if l=='<!-- PAGE_BREAK -->': doc.add_page_break();i+=1;continue
    if l.startswith('# '):
        p=doc.add_paragraph(style=roles['Title']);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        font_run(p.add_run(l[2:]),22,True,NAVY);p.paragraph_format.space_after=Pt(8)
    elif l.startswith('## '):
        p=doc.add_paragraph(style=roles['Heading 1']);font_run(p.add_run(l[3:]),17,True,NAVY)
    elif l.startswith('|'):
        rows=[]
        while i<len(lines) and lines[i].strip().startswith('|'):rows.append(lines[i]);i+=1
        table(rows);continue
    elif l.startswith('**ชื่อบทความ:'):
        p=doc.add_paragraph();p.paragraph_format.space_after=Pt(5);p.paragraph_format.line_spacing=Pt(16)
        while i<len(lines) and lines[i].strip():
            if p.runs:p.add_run().add_break()
            inline(p,lines[i].strip(),14)
            i+=1
        continue
    elif l.startswith('**ตารางที่'):
        p=paragraph(l,13.5,True);p.paragraph_format.keep_with_next=True
    elif l.startswith('[1]'):
        paragraph(l,13.5,True)
    else: paragraph(l)
    i+=1

doc.core_properties.author='นายธีรภัทร ภู่ระย้า'
doc.core_properties.title='รายงานย่อบทความ Publish Time Optimizations'
doc.core_properties.subject='Assignment Paper IV, IEEE VIS 2025, DOI 10.1109/VIS60296.2025.00051'
settings=doc.settings._element;up=OxmlElement('w:updateFields');up.set(qn('w:val'),'true');settings.append(up)
OUT.parent.mkdir(parents=True,exist_ok=True);doc.save(OUT)
evidence={'reference_sha256':hashlib.sha256(REFERENCE.read_bytes()).hexdigest(),'output':OUT.name,'source':SOURCE.name,'section_count':len(doc.sections),'page_size_mm':[round(doc.sections[0].page_width.mm,2),round(doc.sections[0].page_height.mm,2)],'body_chars':len(SOURCE.read_text())}
(ROOT/'.build_v4/report-build.json').write_text(json.dumps(evidence,indent=2,ensure_ascii=False))
print(evidence)
