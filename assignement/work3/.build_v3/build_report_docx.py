import re
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "02_report_v3" / "Assignment_Paper_3_v3_Short_Report_TH.md"
OUTPUT_DOCX = ROOT / "02_report_v3" / "Assignment_Paper_3_v3_Short_Report_TH.docx"

FONT = "Tahoma"
ACCENT = "0F2D4A"
LIGHT_BG = "F1F5F9"

def set_run_font(run, size=None, bold=None, color="000000"):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:cs"), FONT)
    run.font.complex_script = True
    run.font.color.rgb = RGBColor.from_string(color)
    if size:
        run.font.size = Pt(size)
        sz_cs = run._element.get_or_add_rPr().find(qn("w:szCs"))
        if sz_cs is None:
            sz_cs = OxmlElement("w:szCs")
            run._element.get_or_add_rPr().append(sz_cs)
        sz_cs.set(qn("w:val"), str(int(size * 2)))
    if bold is not None:
        run.bold = bold
        if bold:
            run._element.get_or_add_rPr().append(OxmlElement("w:bCs"))

def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)

def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("left", start), ("bottom", bottom), ("right", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))

doc = Document()
for section in doc.sections:
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

md_content = SOURCE.read_text(encoding="utf-8")
lines = md_content.splitlines()

in_table = False
table_rows = []

def flush_table():
    global in_table, table_rows
    if not in_table or not table_rows:
        return
    
    rows_data = []
    for r in table_rows:
        cells = [c.strip() for c in r.split('|')[1:-1]]
        if cells and not all(set(c).issubset({'-', ':', ' '}) for c in cells):
            rows_data.append(cells)
            
    if rows_data:
        num_cols = max(len(r) for r in rows_data)
        tbl = doc.add_table(rows=len(rows_data), cols=num_cols)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, row in enumerate(rows_data):
            is_head = (i == 0)
            for j, val in enumerate(row):
                if j < num_cols:
                    cell = tbl.cell(i, j)
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    set_cell_margins(cell, 90, 110, 90, 110)
                    if is_head:
                        set_cell_shading(cell, ACCENT)
                        p = cell.paragraphs[0]
                        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                        run = p.add_run(val)
                        set_run_font(run, size=9, bold=True, color="FFFFFF")
                    else:
                        if i % 2 == 1:
                            set_cell_shading(cell, "F8FAFC")
                        p = cell.paragraphs[0]
                        val_clean = re.sub(r'\*\*(.+?)\*\*', r'\1', val)
                        run = p.add_run(val_clean)
                        is_bold = "**" in val
                        set_run_font(run, size=8.5, bold=is_bold, color="1E293B")
        doc.add_paragraph()  # Spacing
    in_table = False
    table_rows = []

for line in lines:
    line_s = line.strip()
    if not line_s:
        if in_table:
            flush_table()
        continue

    if line_s.startswith('|') and line_s.endswith('|'):
        in_table = True
        table_rows.append(line_s)
        continue
    elif in_table:
        flush_table()

    if line_s.startswith('# '):
        p = doc.add_paragraph()
        run = p.add_run(line_s[2:])
        set_run_font(run, size=15, bold=True, color=ACCENT)
        p.paragraph_format.space_after = Pt(8)
        continue

    if line_s.startswith('**') and ':**' in line_s:
        p = doc.add_paragraph()
        parts = line_s.split(':**', 1)
        lbl = parts[0].replace('**', '').strip() + ':'
        val = parts[1].strip()
        run_lbl = p.add_run(lbl + " ")
        set_run_font(run_lbl, size=9.5, bold=True, color=ACCENT)
        run_val = p.add_run(re.sub(r'\*\*(.+?)\*\*', r'\1', val))
        set_run_font(run_val, size=9.5, bold=False, color="334155")
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        continue

    if line_s.startswith('## '):
        p = doc.add_paragraph()
        run = p.add_run(line_s[3:])
        set_run_font(run, size=12, bold=True, color=ACCENT)
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        continue

    if line_s.startswith('### '):
        p = doc.add_paragraph()
        run = p.add_run(line_s[4:])
        set_run_font(run, size=10.5, bold=True, color="1E3A8A")
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        continue

    if line_s == '---':
        p = doc.add_paragraph()
        run = p.add_run("―" * 45)
        set_run_font(run, size=8, color="CBD5E1")
        p.paragraph_format.space_after = Pt(6)
        continue

    if line_s.startswith('> '):
        p = doc.add_paragraph()
        run = p.add_run(line_s[2:])
        set_run_font(run, size=9.5, color="334155")
        run.italic = True
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.space_after = Pt(6)
        continue

    if line_s.startswith('- ') or line_s.startswith('* '):
        p = doc.add_paragraph(style='List Bullet')
        text = line_s[2:]
        text_clean = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
        run = p.add_run(text_clean)
        set_run_font(run, size=9.5, color="1E293B")
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.2
        continue

    m_num = re.match(r'^(\d+)\.\s+(.+)$', line_s)
    if m_num:
        p = doc.add_paragraph()
        num_run = p.add_run(f"{m_num.group(1)}. ")
        set_run_font(num_run, size=9.5, bold=True, color=ACCENT)
        text_clean = re.sub(r'\*\*(.+?)\*\*', r'\1', m_num.group(2))
        run = p.add_run(text_clean)
        set_run_font(run, size=9.5, color="1E293B")
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.25
        continue

    p = doc.add_paragraph()
    text_clean = re.sub(r'\*\*(.+?)\*\*', r'\1', line_s)
    run = p.add_run(text_clean)
    set_run_font(run, size=9.5, color="1E293B")
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.25
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

if in_table:
    flush_table()

doc.save(OUTPUT_DOCX)
print(f"Generated DOCX: {OUTPUT_DOCX} ({OUTPUT_DOCX.stat().st_size / 1024:.1f} KB)")
