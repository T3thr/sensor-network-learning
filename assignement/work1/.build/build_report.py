from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "02_report" / "Assignment_Paper_1_IoT-CAD_Short_Report_TH.md"
OUTPUT = ROOT / "02_report" / "Assignment_Paper_1_IoT-CAD_Short_Report_TH.docx"

FONT = "Tahoma"
ACCENT = "173F5F"
LIGHT = "EAF1F5"


def set_run_font(run, size=None, bold=None, color="000000"):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:cs"), FONT)
    run.font.complex_script = True
    run.font.color.rgb = RGBColor.from_string(color)
    if size:
        run.font.size = Pt(size)
        sz_cs = run._element.get_or_add_rPr().find(qn("w:szCs"))
        if sz_cs is None:
            sz_cs = OxmlElement("w:szCs")
            run._element.get_or_add_rPr().append(sz_cs)
        sz_cs.set(qn("w:val"), str(int(size * 2)))
    if bold is not None:
        run.bold = bold
        if bold:
            run._element.get_or_add_rPr().append(OxmlElement("w:bCs"))


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=110, bottom=100, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "5")
        tag.set(qn("w:color"), "D9D9D9")


def add_inline_markup(paragraph, text, size=11.5):
    parts = re.split(r"(\*\*.+?\*\*)", text)
    for part in parts:
        if not part:
            continue
        bold = part.startswith("**") and part.endswith("**")
        value = part[2:-2] if bold else part
        run = paragraph.add_run(value)
        set_run_font(run, size=size, bold=bold)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("หน้า ")
    set_run_font(run, size=9, color="5B6770")
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.72)
section.bottom_margin = Inches(0.68)
section.left_margin = Inches(0.82)
section.right_margin = Inches(0.82)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = FONT
normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
normal._element.rPr.rFonts.set(qn("w:cs"), FONT)
normal.font.size = Pt(11.5)
normal.paragraph_format.line_spacing = 1.18
normal.paragraph_format.space_after = Pt(6)

for style_name, size, before, after in (
    ("Title", 23, 0, 10),
    ("Heading 1", 16, 12, 6),
    ("Heading 2", 13.5, 9, 4),
):
    style = styles[style_name]
    style.font.name = FONT
    style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    style._element.rPr.rFonts.set(qn("w:cs"), FONT)
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

# Remove any theme border from Word's built-in Title style.
title_ppr = styles["Title"]._element.get_or_add_pPr()
title_border = title_ppr.find(qn("w:pBdr"))
if title_border is not None:
    title_ppr.remove(title_border)

footer = section.footer
footer_p = footer.paragraphs[0]
add_page_number(footer_p)

lines = SOURCE.read_text(encoding="utf-8").splitlines()
first = True
metadata_done = False
for raw in lines:
    line = raw.strip()
    if not line:
        continue
    if line.startswith("# "):
        p = doc.add_paragraph(style="Title")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(line[2:])
        set_run_font(run, size=23, bold=True)
        first = False
        continue
    if line.startswith("## "):
        if line == "## 1 ที่มาและปัญหาวิจัย" and not metadata_done:
            doc.add_page_break()
            metadata_done = True
        p = doc.add_paragraph(style="Heading 1")
        add_inline_markup(p, line[3:], size=16)
        continue
    if line.startswith("**") and ":**" in line:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(2)
        add_inline_markup(p, line, size=10.8)
        continue
    if line.startswith("[") and re.match(r"\[\d+\]", line):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.22)
        p.paragraph_format.first_line_indent = Inches(-0.22)
        add_inline_markup(p, line, size=10)
        continue
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Inches(0.3)
    add_inline_markup(p, line, size=11.5)

    if line.startswith("ผู้เขียนเก็บข้อมูลปกติ 8 วัน"):
        table = doc.add_table(rows=1, cols=5)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        widths = [2.05, 1.0, 1.0, 1.0, 1.0]
        headers = ["วิธี", "Precision", "Recall", "F0.5", "F1"]
        data = [
            ["IoT-CAD", "92%", "56%", "81%", "70%"],
            ["LSTM baseline", "64%", "44%", "58%", "52%"],
            ["ConvLSTM baseline", "51%", "95%", "56%", "66%"],
            ["One-Class SVM", "89%", "25%", "60%", "39%"],
        ]
        for idx, cell in enumerate(table.rows[0].cells):
            cell.width = Inches(widths[idx])
            cell.text = headers[idx]
            set_cell_shading(cell, ACCENT)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for run in cell.paragraphs[0].runs:
                set_run_font(run, size=9, bold=True, color="FFFFFF")
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        for ridx, row in enumerate(data):
            cells = table.add_row().cells
            for idx, value in enumerate(row):
                cells[idx].width = Inches(widths[idx])
                cells[idx].text = value
                if ridx % 2 == 1:
                    set_cell_shading(cells[idx], LIGHT)
                set_cell_margins(cells[idx])
                cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                cells[idx].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT if idx == 0 else WD_ALIGN_PARAGRAPH.CENTER
                for run in cells[idx].paragraphs[0].runs:
                    set_run_font(run, size=9.2, bold=(ridx == 0))
        set_table_borders(table)
        caption = doc.add_paragraph("ตารางที่ 1 ผลเปรียบเทียบจาก Table 2 ของบทความต้นฉบับ")
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in caption.runs:
            set_run_font(run, size=9, color="5B6770")

core = doc.core_properties
core.title = "รายงานย่อบทความ IoT-CAD"
core.subject = "Assignment Paper 1 วิชาเครือข่ายเซนเซอร์"
core.author = "[กรอกชื่อ นามสกุล]"
core.keywords = "IoT, Sensor Network, Anomaly Detection, IoT-CAD"

doc.save(OUTPUT)
print(OUTPUT)
