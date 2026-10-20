"""PowerPoint backend: renders the shared deck spec with python-pptx.

Consumes the same ``deck_theme`` elements as ``render_html``: 1 px on the
1280 x 720 spec canvas is 9525 EMU, so the 16:9 stage is 12192000 x 6858000 EMU
and both outputs land on identical coordinates. Charts are emitted as real
PowerPoint chart parts (editable data), not as pictures.
"""
from __future__ import annotations

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_TICK_MARK
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import nsdecls, qn
from pptx.util import Emu, Pt
from lxml import etree

from deck_theme import FONT, H, Run, Slide, W

EMU_PER_PX = 9525

_ALIGN = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}
_ANCHOR = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}
_BASELINE = {"sub": "-25000", "sup": "30000"}


def px(v: float) -> Emu:
    return Emu(int(round(v * EMU_PER_PX)))


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color.lstrip("#").upper())


def _no_outline(shape) -> None:
    shape.line.fill.background()
    shape.shadow.inherit = False


# ------------------------------------------------------------------- text
def _apply_run(run, spec: Run, base_size: float, base_color: str, base_bold: bool,
               base_italic: bool, letter_spacing: float) -> None:
    font = run.font
    font.name = FONT
    font.size = Pt(spec.size or base_size)
    font.bold = spec.bold or base_bold
    font.italic = spec.italic or base_italic
    font.color.rgb = rgb(spec.color or base_color)
    rPr = font._rPr
    if spec.baseline in _BASELINE:
        rPr.set("baseline", _BASELINE[spec.baseline])
    if letter_spacing:
        rPr.set("spc", str(int(round(letter_spacing * 0.75 * 100))))
    # Latin + complex-script (Thai) faces must both be Tahoma or PowerPoint
    # silently substitutes a fallback for the Thai runs.
    for tag in ("a:latin", "a:cs", "a:ea"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", FONT)


def _fill_text_frame(tf, runs: list[Run], p) -> None:
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = _ANCHOR[p["valign"]]

    para = tf.paragraphs[0]
    first = True
    for spec in runs:
        for i, segment in enumerate(spec.text.split("\n")):
            if i > 0:
                para = tf.add_paragraph()
                first = False
            if segment == "" and first and len(runs) == 1:
                continue
            run = para.add_run()
            run.text = segment
            _apply_run(run, spec, p["size"], p["color"], p["bold"], p["italic"],
                       p["letter_spacing"])
    for para in tf.paragraphs:
        para.alignment = _ALIGN[p["align"]]
        para.line_spacing = p["line_height"]
        if not para.runs:  # keep empty lines at the right height
            para.font.size = Pt(p["size"])
            para.font.name = FONT


def add_text(slide, el) -> None:
    box = slide.shapes.add_textbox(px(el.x), px(el.y), px(el.w), px(el.h))
    _fill_text_frame(box.text_frame, el.props["runs"], el.props)


def add_bullets(slide, el) -> None:
    p = el.props
    box = slide.shapes.add_textbox(px(el.x), px(el.y), px(el.w), px(el.h))
    tf = box.text_frame
    cfg = dict(p, valign="top", bold=False, italic=False, align="left", letter_spacing=0)
    flat: list[Run] = []
    for i, runs in enumerate(p["items"]):
        if i:
            flat.append(Run("\n"))
        flat.extend(runs)
    _fill_text_frame(tf, flat, cfg)

    mar = int(round(p["indent"] * 0.75 * 12700))
    for i, para in enumerate(tf.paragraphs):
        pPr = para._p.get_or_add_pPr()
        pPr.set("marL", str(mar))
        pPr.set("indent", str(-mar))
        if i:
            para.space_before = Pt(p["gap"] * 0.75)
        pPr.append(etree.fromstring(
            f'<a:buClr {nsdecls("a")}><a:srgbClr val="{p["bullet_color"].lstrip("#").upper()}"/></a:buClr>'))
        pPr.append(etree.fromstring(f'<a:buFont {nsdecls("a")} typeface="Arial"/>'))
        pPr.append(etree.fromstring(f'<a:buChar {nsdecls("a")} char="{p["glyph"]}"/>'))


# ------------------------------------------------------------------ shapes
def add_rect(slide, el) -> None:
    p = el.props
    geom = MSO_SHAPE.ROUNDED_RECTANGLE if p.get("radius") else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(geom, px(el.x), px(el.y), px(el.w), px(el.h))
    if p.get("radius"):
        # adjustment is a fraction of the shorter side
        shape.adjustments[0] = min(0.5, p["radius"] / min(el.w, el.h))
    if p.get("fill"):
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(p["fill"])
    else:
        shape.fill.background()
    if p.get("stroke") and p.get("stroke_w"):
        shape.line.color.rgb = rgb(p["stroke"])
        shape.line.width = Pt(p["stroke_w"] * 0.75)
    else:
        shape.line.fill.background()
    shape.shadow.inherit = False
    shape.text_frame.text = ""


def add_image(slide, el) -> None:
    slide.shapes.add_picture(el.props["path"], px(el.x), px(el.y), px(el.w), px(el.h))


def add_arrow(slide, el) -> None:
    p = el.props
    head, color = p["head"], p["color"]
    shaft_w = el.w - head
    shaft = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, px(el.x), px(el.y - el.h / 2),
                                   px(shaft_w), px(el.h))
    shaft.fill.solid()
    shaft.fill.fore_color.rgb = rgb(color)
    _no_outline(shaft)

    # Isosceles triangle rotated 90 deg points right; the bounding box keeps its
    # extents, so width/height are swapped relative to the drawn triangle.
    tw, th = head * 1.24, head
    cx, cy = el.x + shaft_w + head / 2, el.y
    tri = slide.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE,
                                 px(cx - tw / 2), px(cy - th / 2), px(tw), px(th))
    tri.rotation = 90
    tri.fill.solid()
    tri.fill.fore_color.rgb = rgb(color)
    _no_outline(tri)


# ------------------------------------------------------------------- chart
def _spPr(tag: str) -> etree._Element:
    return etree.fromstring(
        f'<c:spPr {nsdecls("c", "a")}><a:noFill/><a:ln><a:noFill/></a:ln></c:spPr>')


def add_chart(slide, el, rule_color: str, ink: str, muted: str) -> None:
    p = el.props
    data = CategoryChartData()
    data.categories = p["categories"]
    for ser in p["series"]:
        data.add_series(ser["name"], ser["values"])

    frame = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,
                                   px(el.x), px(el.y), px(el.w), px(el.h), data)
    chart = frame.chart
    chart.has_title = False
    chart.has_legend = False
    chart.font.name = FONT
    chart.font.size = Pt(12)
    chart.font.color.rgb = rgb(ink)

    plot = chart.plots[0]
    plot.gap_width = 127
    plot.vary_by_categories = False
    plot.has_data_labels = True
    labels = plot.data_labels
    labels.position = XL_LABEL_POSITION.OUTSIDE_END
    labels.number_format = '0.0"' + p.get("axis_unit", "") + '"' if p.get("axis_unit") else "0.0"
    labels.number_format_is_linked = False
    labels.font.name = FONT
    labels.font.size = Pt(12)
    labels.font.bold = True

    for ser in chart.series:
        spec = p["series"][0]
        colors = spec.get("colors") or [spec.get("color", "#6D28D9")] * len(spec["values"])
        for i, point in enumerate(ser.points):
            point.format.fill.solid()
            point.format.fill.fore_color.rgb = rgb(colors[i])
            point.format.line.fill.background()
            point.data_label.font.name = FONT
            point.data_label.font.size = Pt(12)
            point.data_label.font.bold = True
            point.data_label.font.color.rgb = rgb(colors[i])

    va = chart.value_axis
    va.minimum_scale = 0
    va.maximum_scale = p["ymax"]
    va.major_unit = p["ystep"]
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = rgb(rule_color)
    va.major_gridlines.format.line.width = Pt(0.75)
    va.major_tick_mark = XL_TICK_MARK.NONE
    va.minor_tick_mark = XL_TICK_MARK.NONE
    va.format.line.fill.background()
    va.tick_labels.font.name = FONT
    va.tick_labels.font.size = Pt(12)
    va.tick_labels.font.color.rgb = rgb(muted)
    va.tick_labels.number_format = "0"
    va.tick_labels.number_format_is_linked = False

    ca = chart.category_axis
    ca.has_major_gridlines = False
    ca.major_tick_mark = XL_TICK_MARK.NONE
    ca.minor_tick_mark = XL_TICK_MARK.NONE
    ca.format.line.color.rgb = rgb("#94A3B8")
    ca.format.line.width = Pt(0.75)
    ca.tick_labels.font.name = FONT
    ca.tick_labels.font.size = Pt(12)
    ca.tick_labels.font.color.rgb = rgb(ink)

    space = chart._chartSpace
    space.find(qn("c:chart")).addnext(_spPr("c:spPr"))
    space.find(".//" + qn("c:plotArea")).append(_spPr("c:spPr"))


# ------------------------------------------------------------------- deck
def render(slides: list[Slide], out_path) -> None:
    from deck_theme import INK, INK_MUTED, RULE

    prs = Presentation()
    prs.slide_width = px(W)
    prs.slide_height = px(H)
    blank = prs.slide_layouts[6]

    for spec in slides:
        slide = prs.slides.add_slide(blank)
        bg = slide.background
        bg.fill.solid()
        bg.fill.fore_color.rgb = rgb(spec.background)

        for el in spec.elements:
            if el.kind == "rect":
                add_rect(slide, el)
            elif el.kind == "text":
                add_text(slide, el)
            elif el.kind == "bullets":
                add_bullets(slide, el)
            elif el.kind == "image":
                add_image(slide, el)
            elif el.kind == "arrow":
                add_arrow(slide, el)
            elif el.kind == "chart":
                add_chart(slide, el, RULE, INK, INK_MUTED)
            else:
                raise ValueError(f"unknown element kind: {el.kind}")

        if spec.notes:
            slide.notes_slide.notes_text_frame.text = spec.notes

    prs.save(str(out_path))
